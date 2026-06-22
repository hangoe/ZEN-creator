"""
Read/write helpers for the technology parametrization workbooks under
`input_data/Parametrization/` (e.g. `process_parametrization.xlsx`,
`heat_tech_parametrization.xlsx`).

Each workbook has a sheet where the first column ("parameter") holds row
labels matching `attributes.json` keys (with `conversion_factor:<carrier>`
rows for individual conversion-factor entries), and one column per
technology holding that technology's value for each parameter.

These helpers let `compute_params.py`:
  1. write freshly computed conversion factors into the workbook, and
  2. read a technology's column back out and apply it as overrides on top
     of a code-generated `attributes.json` dict, so the workbook is the
     final source of truth for default values.
"""

import json
import pathlib

import openpyxl

# Row labels (column "parameter") that hold semicolon-separated carrier lists
# rather than a single scalar value.
LIST_FIELDS = ("reference_carrier", "input_carrier", "output_carrier")

CONVERSION_FACTOR_PREFIX = "conversion_factor:"

# Header columns that are not technologies.
META_COLUMNS = {"parameter", "unit", "source", "comment"}

# Technology columns prefixed with this are comparison-only (e.g. existing
# Crystal Ball techs shown for reference) and must never be written out.
COMPARISON_PREFIX = "XX_"


def _label_to_row(ws) -> dict[str, int]:
    """Map each non-empty value in column A (row >= 2) to its row number."""
    return {
        ws.cell(row=row, column=1).value: row
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=1).value is not None
    }


def _column_index(ws, column_header: str) -> int:
    for cell in ws[1]:
        if cell.value == column_header:
            return cell.column
    raise KeyError(f"Column {column_header!r} not found in sheet {ws.title!r}")


def load_param_column(xlsx_path: pathlib.Path, sheet_name: str, column_header: str) -> dict[str, object]:
    """Read one technology's column into {parameter_row_label: value}.

    Blank cells are included as `None`; callers should skip these when
    applying overrides so that blanks fall back to code-computed defaults.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet_name]
    col = _column_index(ws, column_header)
    return {label: ws.cell(row=row, column=col).value for label, row in _label_to_row(ws).items()}


def write_param_column(xlsx_path: pathlib.Path, sheet_name: str, column_header: str, values: dict[str, object]) -> None:
    """Write `values` (parameter_row_label -> value) into one technology's column.

    Existing rows are updated in place; labels with no matching row are
    appended as new rows at the bottom of the sheet.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet_name]
    col = _column_index(ws, column_header)
    label_to_row = _label_to_row(ws)
    next_row = ws.max_row + 1
    for label, value in values.items():
        row = label_to_row.get(label)
        if row is None:
            row = next_row
            ws.cell(row=row, column=1, value=label)
            next_row += 1
            label_to_row[label] = row
        ws.cell(row=row, column=col, value=value)
    wb.save(xlsx_path)


# Parameter rows whose units already match between a Crystal Ball
# `attributes.json` and `heat_tech_parametrization.xlsx` (years, Euro/MWh,
# Euro/kW, "1", kilotons/GWh, GW, GW/GW), so the value can be copied across
# as-is. "input_carrier" and "conversion_factor:input_carrier" are resolved
# specially (see `_crystal_ball_param_value`).
#
# reference_carrier and output_carrier are deliberately excluded: Crystal
# Ball's techs reference/output "heat", but the _industry techs use
# "heat_low_temp_industry" in the Excel (remapped to heat_industry_0_100 /
# heat_industry_100_200 by compute_params.py when writing the JSON).
HEAT_TECH_CRYSTAL_BALL_PARAMS = (
    "min_load",
    "max_load",
    "lifetime",
    "opex_specific_variable",
    "carbon_intensity_technology",
    "construction_time",
    "capacity_investment_existing",
    "opex_specific_fixed",
    "max_diffusion_rate",
    "capex_specific_conversion",
    "input_carrier",
    "conversion_factor:input_carrier",
)

# Crystal Ball technology name -> heat_tech_parametrization.xlsx column header
# to write its parameters into. heat_pump_industry is intentionally not
# included: it is a custom parametrization for high-temperature industrial
# heat pumps, distinct from Crystal Ball's residential heat_pump (which is
# kept for reference in the XX_heat_pump column).
CRYSTAL_BALL_HEAT_TECH_COLUMNS = {
    "electrode_boiler": "electrode_boiler_industry",
    "natural_gas_boiler": "natural_gas_boiler_industry",
    "biomass_boiler": "biomass_boiler_industry",
    "heat_pump": "XX_heat_pump",
}


def update_heat_techs_from_crystal_ball(
    xlsx_path: pathlib.Path,
    sheet_name: str,
    crystal_ball_dir: pathlib.Path,
    tech_map: dict[str, str] = CRYSTAL_BALL_HEAT_TECH_COLUMNS,
) -> None:
    """Copy heat-technology parameters from a Crystal Ball dataset into
    `heat_tech_parametrization.xlsx`.

    For each `crystal_ball_tech -> column_header` pair in `tech_map`
    (default `CRYSTAL_BALL_HEAT_TECH_COLUMNS`), reads
    `crystal_ball_dir/<crystal_ball_tech>/attributes.json` and writes
    `HEAT_TECH_CRYSTAL_BALL_PARAMS` into `column_header`.
    """
    for cb_tech, column_header in tech_map.items():
        with open(pathlib.Path(crystal_ball_dir) / cb_tech / "attributes.json") as f:
            data = json.load(f)

        values = {param: _crystal_ball_param_value(data, param) for param in HEAT_TECH_CRYSTAL_BALL_PARAMS}
        write_param_column(xlsx_path, sheet_name, column_header, values)


def _crystal_ball_param_value(data: dict, param: str) -> object:
    """Resolve `param` (a `HEAT_TECH_CRYSTAL_BALL_PARAMS` entry) from a
    Crystal Ball `attributes.json` dict.

    - "input_carrier" is joined into the "; "-separated string format used
      by the worksheet's `LIST_FIELDS` cells.
    - "conversion_factor:input_carrier" is the conversion factor of the
      technology's single input carrier.
    - All other params are top-level `attributes.json` keys.
    """
    if param == "input_carrier":
        return "; ".join(data["input_carrier"]["default_value"])
    if param == "conversion_factor:input_carrier":
        [input_carrier] = data["input_carrier"]["default_value"]
        return next(entry[input_carrier]["default_value"] for entry in data["conversion_factor"] if input_carrier in entry)
    return data[param]["default_value"]


def tech_columns(xlsx_path: pathlib.Path, sheet_name: str) -> list[str]:
    """List the technology column headers of a parametrization sheet.

    Excludes the metadata columns ("parameter", "unit", "source", "comment")
    and comparison-only columns prefixed with `COMPARISON_PREFIX` (e.g.
    "XX_heat_pump"), which are shown for reference only and must not be
    written out as conversion technologies.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet_name]
    return [
        cell.value for cell in ws[1]
        if cell.value not in META_COLUMNS and not str(cell.value).startswith(COMPARISON_PREFIX)
    ]


def build_tech_from_table(xlsx_path: pathlib.Path, sheet_name: str, column_header: str) -> dict:
    """Build a complete `attributes.json` dict for a technology purely from
    its parametrization table column.

    Every row in the sheet becomes an attribute: scalar rows become
    `{"default_value": value, "unit": <column B>}`, `LIST_FIELDS` rows
    become semicolon-split carrier lists, and `conversion_factor:<suffix>`
    rows become entries in `data["conversion_factor"]`, keyed by the
    carrier name in the `<suffix>` row (e.g. "conversion_factor:input_carrier"
    -> the technology's `input_carrier` value).

    Raises `ValueError` if any cell in the column is blank, since (unlike
    `apply_excel_overrides`) there is no code-computed fallback here - the
    table must be complete.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet_name]
    col = _column_index(ws, column_header)

    rows = {label: (ws.cell(row=row, column=col).value, ws.cell(row=row, column=2).value) for label, row in _label_to_row(ws).items()}

    missing = [label for label, (value, _) in rows.items() if value is None]
    if missing:
        raise ValueError(f"{xlsx_path.name}!{sheet_name}: column {column_header!r} is missing values for {missing}")

    list_field_values = {label: value for label, (value, _) in rows.items() if label in LIST_FIELDS}

    data: dict[str, object] = {}
    conversion_factor = []
    for label, (value, unit) in rows.items():
        if label.startswith(CONVERSION_FACTOR_PREFIX):
            suffix = label[len(CONVERSION_FACTOR_PREFIX):]
            carrier = str(list_field_values.get(suffix, suffix)).strip()
            conversion_factor.append({carrier: {"default_value": value, "unit": unit}})
        elif label in LIST_FIELDS:
            data[label] = {"default_value": [v.strip() for v in str(value).split(";")]}
        else:
            data[label] = {"default_value": value, "unit": unit}

    data["conversion_factor"] = conversion_factor
    return data


def apply_excel_overrides(data: dict, overrides: dict[str, object], conversion_factor_map: dict[str, str] | None = None) -> dict:
    """Overwrite `default_value`s in a code-generated `attributes.json` dict
    with non-blank values from `overrides` (as returned by `load_param_column`).

    - `conversion_factor:<label>` rows update the matching entry in
      `data["conversion_factor"]`. `conversion_factor_map` maps the row
      label (e.g. "conversion_factor:natural_gas") to the actual carrier name
      used in `data["conversion_factor"]` (e.g. "natural_gas"); rows not
      in the map default to the carrier name after the prefix.
    - `reference_carrier`, `input_carrier`, `output_carrier` rows are parsed
      as semicolon-separated carrier lists.
    - All other rows overwrite `data[label]["default_value"]` if `label` is
      a key of `data`.
    - Blank (`None`) values are skipped, leaving the code-computed default.
    """
    conversion_factor_map = conversion_factor_map or {}
    for label, value in overrides.items():
        if value is None or label is None:
            continue

        if label.startswith(CONVERSION_FACTOR_PREFIX):
            carrier = conversion_factor_map.get(label, label[len(CONVERSION_FACTOR_PREFIX):])
            for entry in data["conversion_factor"]:
                if carrier in entry:
                    entry[carrier]["default_value"] = value
                    break
            continue

        if label not in data:
            continue

        if label in LIST_FIELDS:
            data[label]["default_value"] = [v.strip() for v in str(value).split(";")]
        else:
            data[label]["default_value"] = value

    return data
