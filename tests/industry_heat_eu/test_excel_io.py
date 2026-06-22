import json

import openpyxl
import pytest

from zen_creator.industry_heat_eu.excel_io import (
    apply_excel_overrides,
    build_tech_from_table,
    load_param_column,
    tech_columns,
    update_heat_techs_from_crystal_ball,
    write_param_column,
)


@pytest.fixture
def sample_xlsx(tmp_path):
    path = tmp_path / "params.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "techs"
    ws.append(["parameter", "unit", "glass_production"])
    ws.append(["lifetime", "1", 15])
    ws.append(["capex_specific_conversion", "Euro/(tonproduct/h)", None])
    ws.append(["input_carrier", "-", "fuel_for_glass; heat_low_temp_industry; electricity"])
    ws.append(["conversion_factor:fuel_for_X", "GW/(tonproduct/hour)", None])
    wb.save(path)
    return path


def test_load_param_column_reads_values_and_blanks(sample_xlsx):
    values = load_param_column(sample_xlsx, "techs", "glass_production")
    assert values["lifetime"] == 15
    assert values["capex_specific_conversion"] is None
    assert values["input_carrier"] == "fuel_for_glass; heat_low_temp_industry; electricity"


def test_write_param_column_updates_existing_and_appends_new_rows(sample_xlsx):
    write_param_column(sample_xlsx, "techs", "glass_production", {
        "capex_specific_conversion": 5350000,
        "conversion_factor:fuel_for_X": 0.00113,
        "max_diffusion_rate": 0.13,
    })

    values = load_param_column(sample_xlsx, "techs", "glass_production")
    assert values["capex_specific_conversion"] == 5350000
    assert values["conversion_factor:fuel_for_X"] == 0.00113
    assert values["max_diffusion_rate"] == 0.13
    # untouched rows keep their value
    assert values["lifetime"] == 15


def test_apply_excel_overrides_updates_scalars_lists_and_conversion_factors():
    data = {
        "lifetime": {"default_value": 99, "unit": "1"},
        "capex_specific_conversion": {"default_value": 1.0, "unit": "Euro/(tonproduct/h)"},
        "input_carrier": {"default_value": ["placeholder"]},
        "conversion_factor": [
            {"fuel_for_glass": {"default_value": 0.0, "unit": "GW/(tonproduct/hour)"}},
            {"electricity": {"default_value": 0.0, "unit": "GW/(tonproduct/hour)"}},
        ],
    }
    overrides = {
        "lifetime": 15,
        "capex_specific_conversion": 5350000,
        "input_carrier": "fuel_for_glass; heat_low_temp_industry; electricity",
        "conversion_factor:fuel_for_X": 0.00113,
        "carbon_intensity_technology": None,  # not in data, must be ignored
    }

    result = apply_excel_overrides(data, overrides, conversion_factor_map={"conversion_factor:fuel_for_X": "fuel_for_glass"})

    assert result["lifetime"]["default_value"] == 15
    assert result["capex_specific_conversion"]["default_value"] == 5350000
    assert result["input_carrier"]["default_value"] == ["fuel_for_glass", "heat_low_temp_industry", "electricity"]
    factors = {list(entry)[0]: entry for entry in result["conversion_factor"]}
    assert factors["fuel_for_glass"]["fuel_for_glass"]["default_value"] == 0.00113
    assert factors["electricity"]["electricity"]["default_value"] == 0.0  # untouched
    assert "carbon_intensity_technology" not in result


def test_apply_excel_overrides_skips_blank_values():
    data = {"capex_specific_conversion": {"default_value": 1.0, "unit": "Euro/(tonproduct/h)"}}
    result = apply_excel_overrides(data, {"capex_specific_conversion": None})
    assert result["capex_specific_conversion"]["default_value"] == 1.0


def test_apply_excel_overrides_conversion_factor_without_map_uses_suffix():
    data = {"conversion_factor": [{"hard_coal": {"default_value": 0.0, "unit": "GW/GW"}}]}
    result = apply_excel_overrides(data, {"conversion_factor:hard_coal": 0.174828})
    assert result["conversion_factor"][0]["hard_coal"]["default_value"] == 0.174828


@pytest.fixture
def heat_techs_xlsx(tmp_path):
    path = tmp_path / "heat_techs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "heat_techs"
    ws.append(["parameter", "unit", "heat_pump_industry", "XX_heat_pump", "source", "comment"])
    ws.append(["lifetime", "1 (years)", 15, 19, "src", None])
    ws.append(["capex_specific_conversion", "Euro/kW", 900, 930, "src", None])
    ws.append(["reference_carrier", "-", "heat_low_temp_industry", "heat", None, None])
    ws.append(["input_carrier", "-", "electricity", "electricity", None, None])
    ws.append(["output_carrier", "-", "heat_low_temp_industry", "heat", None, None])
    ws.append(["conversion_factor:input_carrier", "GW/GW", 0.33, 0.33, "src", None])
    wb.save(path)
    return path


def test_tech_columns_excludes_meta_and_comparison_columns(heat_techs_xlsx):
    assert tech_columns(heat_techs_xlsx, "heat_techs") == ["heat_pump_industry"]


@pytest.fixture
def crystal_ball_dir(tmp_path):
    cb_dir = tmp_path / "Crystal_Ball" / "set_technologies" / "set_conversion_technologies"

    def write_tech(name, input_carrier, conversion_factor):
        tech_dir = cb_dir / name
        tech_dir.mkdir(parents=True)
        (tech_dir / "attributes.json").write_text(json.dumps({
            "min_load": {"default_value": 0, "unit": "1"},
            "max_load": {"default_value": 1, "unit": "1"},
            "lifetime": {"default_value": 21, "unit": "1"},
            "opex_specific_variable": {"default_value": 0.5, "unit": "Euro/MWh"},
            "carbon_intensity_technology": {"default_value": 0, "unit": "kilotons/GWh"},
            "construction_time": {"default_value": 0, "unit": "1"},
            "capacity_investment_existing": {"default_value": 0, "unit": "GW"},
            "opex_specific_fixed": {"default_value": 17.36, "unit": "Euro/kW"},
            "max_diffusion_rate": {"default_value": 0.29, "unit": "1"},
            "capex_specific_conversion": {"default_value": 487.19, "unit": "Euro/kW"},
            "reference_carrier": {"default_value": ["heat"]},
            "input_carrier": {"default_value": [input_carrier]},
            "output_carrier": {"default_value": ["heat"]},
            "conversion_factor": [{input_carrier: {"default_value": conversion_factor, "unit": "GW/GW"}}],
        }))

    write_tech("natural_gas_boiler", "natural_gas", 1.005)
    write_tech("biomass_boiler", "biomass", 1.198)
    write_tech("electrode_boiler", "electricity", 1.0)
    write_tech("heat_pump", "electricity", 0.3306)
    return cb_dir


@pytest.fixture
def heat_techs_for_cb_xlsx(tmp_path):
    path = tmp_path / "heat_techs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "heat_techs"
    ws.append(["parameter", "unit", "natural_gas_boiler_industry", "biomass_boiler_industry", "electrode_boiler_industry", "XX_heat_pump"])
    ws.append(["min_load", "1", None, None, None, None])
    ws.append(["max_load", "1", None, None, None, None])
    ws.append(["lifetime", "1 (years)", None, None, None, None])
    ws.append(["opex_specific_variable", "Euro/MWh", None, None, None, None])
    ws.append(["carbon_intensity_technology", "kilotons/GWh", None, None, None, None])
    ws.append(["construction_time", "1", None, None, None, None])
    ws.append(["capacity_investment_existing", "GW", None, None, None, None])
    ws.append(["opex_specific_fixed", "Euro/kW", None, None, None, None])
    ws.append(["max_diffusion_rate", "1", None, None, None, None])
    ws.append(["capex_specific_conversion", "Euro/kW", None, None, None, None])
    ws.append(["reference_carrier", "-", "heat_low_temp_industry", "heat_low_temp_industry", "heat_low_temp_industry", "heat"])
    ws.append(["input_carrier", "-", None, None, None, None])
    ws.append(["output_carrier", "-", "heat_low_temp_industry", "heat_low_temp_industry", "heat_low_temp_industry", "heat"])
    ws.append(["conversion_factor:input_carrier", "GW/GW", None, None, None, None])
    wb.save(path)
    return path


def test_update_heat_techs_from_crystal_ball_copies_matching_params(heat_techs_for_cb_xlsx, crystal_ball_dir):
    update_heat_techs_from_crystal_ball(heat_techs_for_cb_xlsx, "heat_techs", crystal_ball_dir)

    ng = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "natural_gas_boiler_industry")
    assert ng["min_load"] == 0
    assert ng["max_load"] == 1
    assert ng["lifetime"] == 21
    assert ng["opex_specific_variable"] == 0.5
    assert ng["carbon_intensity_technology"] == 0
    assert ng["construction_time"] == 0
    assert ng["capacity_investment_existing"] == 0
    assert ng["opex_specific_fixed"] == 17.36
    assert ng["max_diffusion_rate"] == 0.29
    assert ng["capex_specific_conversion"] == 487.19
    assert ng["input_carrier"] == "natural_gas"
    assert ng["conversion_factor:input_carrier"] == 1.005

    biomass = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "biomass_boiler_industry")
    assert biomass["input_carrier"] == "biomass"
    assert biomass["conversion_factor:input_carrier"] == 1.198

    electrode = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "electrode_boiler_industry")
    assert electrode["input_carrier"] == "electricity"
    assert electrode["conversion_factor:input_carrier"] == 1.0

    xx_heat_pump = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "XX_heat_pump")
    assert xx_heat_pump["input_carrier"] == "electricity"
    assert xx_heat_pump["conversion_factor:input_carrier"] == 0.3306


def test_update_heat_techs_from_crystal_ball_does_not_overwrite_reference_or_output_carrier(heat_techs_for_cb_xlsx, crystal_ball_dir):
    update_heat_techs_from_crystal_ball(heat_techs_for_cb_xlsx, "heat_techs", crystal_ball_dir)

    ng = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "natural_gas_boiler_industry")
    assert ng["reference_carrier"] == "heat_low_temp_industry"
    assert ng["output_carrier"] == "heat_low_temp_industry"

    # XX_heat_pump's "heat" reference/output carriers (Crystal Ball's actual
    # values) are left untouched too.
    xx_heat_pump = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "XX_heat_pump")
    assert xx_heat_pump["reference_carrier"] == "heat"
    assert xx_heat_pump["output_carrier"] == "heat"


def test_update_heat_techs_from_crystal_ball_skips_heat_pump_industry(heat_techs_for_cb_xlsx, crystal_ball_dir):
    wb = openpyxl.load_workbook(heat_techs_for_cb_xlsx)
    ws = wb["heat_techs"]
    ws.cell(row=1, column=7, value="heat_pump_industry")
    wb.save(heat_techs_for_cb_xlsx)
    # custom lifetime, distinct from Crystal Ball heat_pump's 21
    write_param_column(heat_techs_for_cb_xlsx, "heat_techs", "heat_pump_industry", {"lifetime": 15})

    update_heat_techs_from_crystal_ball(heat_techs_for_cb_xlsx, "heat_techs", crystal_ball_dir)

    heat_pump_industry = load_param_column(heat_techs_for_cb_xlsx, "heat_techs", "heat_pump_industry")
    assert heat_pump_industry["lifetime"] == 15


def test_build_tech_from_table_builds_full_attributes(heat_techs_xlsx):
    data = build_tech_from_table(heat_techs_xlsx, "heat_techs", "heat_pump_industry")

    assert data["lifetime"] == {"default_value": 15, "unit": "1 (years)"}
    assert data["capex_specific_conversion"] == {"default_value": 900, "unit": "Euro/kW"}
    assert data["reference_carrier"]["default_value"] == ["heat_low_temp_industry"]
    assert data["input_carrier"]["default_value"] == ["electricity"]
    assert data["conversion_factor"] == [{"electricity": {"default_value": 0.33, "unit": "GW/GW"}}]


def test_build_tech_from_table_raises_on_missing_value(heat_techs_xlsx):
    wb = openpyxl.load_workbook(heat_techs_xlsx)
    ws = wb["heat_techs"]
    ws["C2"] = None  # blank out lifetime for heat_pump_industry
    wb.save(heat_techs_xlsx)

    with pytest.raises(ValueError, match="lifetime"):
        build_tech_from_table(heat_techs_xlsx, "heat_techs", "heat_pump_industry")
